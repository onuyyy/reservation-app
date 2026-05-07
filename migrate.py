"""
사용법:
  python migrate.py                          # 현재 폴더의 모든 .xlsx
  python migrate.py 26-4.xlsx 26-5.xlsx     # 파일 직접 지정
  python migrate.py /path/to/folder         # 폴더 안의 모든 .xlsx

  --reset 플래그: 해당 파일의 연월 예약을 먼저 삭제 후 재삽입 (venues/clients 유지)
  python migrate.py --reset 26-4.xlsx
"""
import sys
import re
import os
import glob
from datetime import date
import openpyxl
from sqlalchemy import text, extract
from database import SessionLocal, init_db, Venue, Client, Reservation


# ── 인자 파싱 ───────────────────────────────────────────────
def parse_args(argv: list[str]) -> tuple[bool, list[str]]:
    reset = "--reset" in argv
    paths = [a for a in argv if a != "--reset"]
    return reset, paths


# ── 파일 목록 결정 ──────────────────────────────────────────
def resolve_paths(args: list[str]) -> list[str]:
    if not args:
        files = sorted(glob.glob("*.xlsx"))
        if not files:
            print("[ERROR] 현재 폴더에 .xlsx 파일이 없습니다.")
            sys.exit(1)
        return files

    files = []
    for arg in args:
        if os.path.isdir(arg):
            found = sorted(glob.glob(os.path.join(arg, "*.xlsx")))
            if not found:
                print(f"[WARN] 폴더에 .xlsx 없음: {arg}")
            files.extend(found)
        elif os.path.isfile(arg):
            files.append(arg)
        else:
            print(f"[WARN] 파일/폴더 없음, 건너뜀: {arg}")
    return files


# ── 파일명에서 연도/월 추출 (예: 26-4.xlsx → 2026, 4) ──────
def year_month_from_filename(path: str) -> tuple[int, int] | tuple[None, None]:
    name = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r'(\d{2,4})[-_](\d{1,2})', name)
    if m:
        y = int(m.group(1))
        if y < 100:
            y += 2000
        return y, int(m.group(2))
    return None, None


# ── 변환 헬퍼 ───────────────────────────────────────────────
def to_float(v):
    try:
        return float(v) if v is not None else 0.0
    except (ValueError, TypeError):
        return 0.0

def to_int(v):
    try:
        return int(v) if v is not None else 0
    except (ValueError, TypeError):
        return 0

def to_str(v):
    return str(v).strip() if v is not None else ""


# ── 날짜 파싱 ───────────────────────────────────────────────
def parse_date(title: str, sheet_name: str, default_year: int, default_month: int) -> date | None:
    m = re.search(r'(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})', str(title))
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
    day = day_from_sheetname(sheet_name)
    if day and default_year and default_month:
        try:
            return date(default_year, default_month, day)
        except ValueError:
            pass
    return None


# 시트명에서 날짜용 일(day) 숫자 추출
# '8' → 8,  '1-8' → 8,  '합계' → None
def day_from_sheetname(sname: str) -> int | None:
    if sname.isdigit():
        return int(sname)
    m = re.fullmatch(r'\d+-(\d+)', sname)
    if m:
        return int(m.group(1))
    return None


# ── 엑셀 파일 파싱 ──────────────────────────────────────────
def parse_excel(path: str) -> tuple[dict, dict, list]:
    default_year, default_month = year_month_from_filename(path)
    wb = openpyxl.load_workbook(path, data_only=True)

    venue_map, client_map, rows = {}, {}, []

    for sname in wb.sheetnames:
        if day_from_sheetname(sname) is None:
            continue

        ws = wb[sname]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue

        event_date = parse_date(all_rows[0][0], sname, default_year, default_month)
        if event_date is None:
            print(f"  [SKIP] 시트 '{sname}': 날짜 파싱 실패")
            continue

        header_idx = next(
            (i for i, row in enumerate(all_rows) if row[0] and to_str(row[0]) == "단체명"),
            None
        )
        if header_idx is None:
            continue

        for row in all_rows[header_idx + 1:]:
            a = to_str(row[0])
            if not a or a in ("합 계", "합계", "단체명"):
                continue

            group_name = a
            event_name = to_str(row[1])
            if event_name:
                venue_map[event_name] = True
            client_map[group_name] = True

            rows.append({
                "event_date":     event_date,
                "group_name":     group_name,
                "event_name":     event_name,
                "vendor":         event_name,
                "sale_price":     to_float(row[2]),
                "dc_sale_price":  to_float(row[3]),
                "headcount":      to_int(row[4]),
                "dc_amount":      to_float(row[5]),
                "deposit":        to_float(row[6]),
                "actual_sale":    to_float(row[7]),
                "phone":          to_str(row[8]),
                "bank":           to_str(row[9]),
                "account_number": to_str(row[10]),
                "account_holder": to_str(row[11]),
                "manager":        to_str(row[12]),
                "memo":           "",
            })

    return venue_map, client_map, rows


# ── DB 스키마 보정 ──────────────────────────────────────────
def ensure_schema(db):
    conn = db.connection()
    cols = [row[1] for row in conn.execute(text("PRAGMA table_info(reservations)"))]
    added = []
    if "venue_id" not in cols:
        conn.execute(text("ALTER TABLE reservations ADD COLUMN venue_id INTEGER REFERENCES venues(id)"))
        added.append("venue_id")
    if "client_id" not in cols:
        conn.execute(text("ALTER TABLE reservations ADD COLUMN client_id INTEGER REFERENCES clients(id)"))
        added.append("client_id")
    if added:
        db.commit()
        print(f"  스키마 추가: {', '.join(added)}")


# ── 연월 리셋 ───────────────────────────────────────────────
def reset_month(db, year: int, month: int):
    deleted = db.query(Reservation).filter(
        extract("year",  Reservation.event_date) == year,
        extract("month", Reservation.event_date) == month,
    ).delete(synchronize_session=False)
    db.commit()
    print(f"  --reset: {year}년 {month}월 예약 {deleted}건 삭제")


# ── DB 저장 ─────────────────────────────────────────────────
def save_to_db(db, venue_map: dict, client_map: dict, rows: list) -> tuple[int, int]:
    # venues upsert
    for name in venue_map:
        if not db.query(Venue).filter(Venue.name == name).first():
            db.add(Venue(name=name))
    db.commit()

    # clients upsert
    for name in client_map:
        if not db.query(Client).filter(Client.name == name).first():
            db.add(Client(name=name))
    db.commit()

    venues_by_name  = {v.name: v.id for v in db.query(Venue).all()}
    clients_by_name = {c.name: c.id for c in db.query(Client).all()}

    inserted = skipped = 0
    for row in rows:
        try:
            exists = db.query(Reservation).filter(
                Reservation.event_date == row["event_date"],
                Reservation.group_name == row["group_name"],
                Reservation.actual_sale == row["actual_sale"],
            ).first()
            if exists:
                skipped += 1
                continue
            db.add(Reservation(
                **row,
                venue_id=venues_by_name.get(row["event_name"]),
                client_id=clients_by_name.get(row["group_name"]),
            ))
            db.commit()
            inserted += 1
        except Exception as e:
            db.rollback()
            print(f"  [ERROR] {row['event_date']} {row['group_name']}: {e}")
            skipped += 1

    return inserted, skipped


# ── 메인 ────────────────────────────────────────────────────
def main():
    do_reset, path_args = parse_args(sys.argv[1:])
    files = resolve_paths(path_args)
    print(f"처리할 파일 {len(files)}개: {[os.path.basename(f) for f in files]}")
    if do_reset:
        print("모드: --reset (해당 연월 예약 삭제 후 재삽입)\n")
    else:
        print("모드: 중복 skip (기존 데이터 유지)\n")

    init_db()
    db = SessionLocal()
    ensure_schema(db)

    total_inserted = total_skipped = 0
    all_venues, all_clients = {}, {}

    for path in files:
        print(f"▶ {os.path.basename(path)}")
        try:
            venue_map, client_map, rows = parse_excel(path)
            print(f"  파싱: venue {len(venue_map)}개, client {len(client_map)}개, 예약 {len(rows)}건")

            if do_reset and rows:
                # 파일에 포함된 연월 조합을 모두 삭제
                year_months = {(r["event_date"].year, r["event_date"].month) for r in rows}
                for y, m in sorted(year_months):
                    reset_month(db, y, m)

            inserted, skipped = save_to_db(db, venue_map, client_map, rows)
            print(f"  저장: 신규 {inserted}건, 중복 skip {skipped}건")
            total_inserted += inserted
            total_skipped  += skipped
            all_venues.update(venue_map)
            all_clients.update(client_map)
        except Exception as e:
            print(f"  [ERROR] 파일 처리 실패: {e}")
        print()

    total_venues  = db.query(Venue).count()
    total_clients = db.query(Client).count()
    total_res     = db.query(Reservation).count()
    db.close()

    print("=" * 40)
    print("전체 완료")
    print(f"  행사장(venues)    : {total_venues}개")
    print(f"  고객(clients)     : {total_clients}개")
    print(f"  예약(reservations): {total_res}건 (신규 {total_inserted}, skip {total_skipped})")
    print("=" * 40)


if __name__ == "__main__":
    main()
