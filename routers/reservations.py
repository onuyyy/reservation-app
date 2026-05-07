from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import extract, func
from typing import Optional
import io
from database import get_db, Reservation
from schemas import ReservationCreate, ReservationUpdate, ReservationOut

router = APIRouter(prefix="/api/reservations", tags=["reservations"])

@router.get("/")
def get_reservations(year: Optional[int]=None, month: Optional[int]=None, vendor: Optional[str]=None, keyword: Optional[str]=None, db: Session=Depends(get_db)):
    q = db.query(Reservation)
    if year: q = q.filter(extract("year", Reservation.event_date)==year)
    if month: q = q.filter(extract("month", Reservation.event_date)==month)
    if vendor: q = q.filter(Reservation.vendor==vendor)
    if keyword: q = q.filter(Reservation.group_name.contains(keyword)|Reservation.event_name.contains(keyword)|Reservation.manager.contains(keyword))
    return q.order_by(Reservation.event_date.desc()).all()

@router.post("/", response_model=ReservationOut)
def create_reservation(data: ReservationCreate, db: Session=Depends(get_db)):
    r = Reservation(**data.model_dump())
    db.add(r); db.commit(); db.refresh(r)
    return r

@router.put("/{rid}", response_model=ReservationOut)
def update_reservation(rid: int, data: ReservationUpdate, db: Session=Depends(get_db)):
    r = db.query(Reservation).filter(Reservation.id==rid).first()
    if not r: raise HTTPException(status_code=404, detail="Not found")
    for k,v in data.model_dump().items(): setattr(r,k,v)
    db.commit(); db.refresh(r)
    return r

@router.delete("/{rid}")
def delete_reservation(rid: int, db: Session=Depends(get_db)):
    r = db.query(Reservation).filter(Reservation.id==rid).first()
    if not r: raise HTTPException(status_code=404, detail="Not found")
    db.delete(r); db.commit()
    return {"ok": True}

@router.get("/stats/monthly")
def monthly_stats(year: int, db: Session=Depends(get_db)):
    rows = db.query(extract("month",Reservation.event_date).label("month"),func.sum(Reservation.headcount).label("headcount"),func.sum(Reservation.actual_sale).label("actual_sale"),func.sum(Reservation.deposit).label("deposit")).filter(extract("year",Reservation.event_date)==year).group_by("month").order_by("month").all()
    return [{"month":int(r.month),"headcount":r.headcount or 0,"actual_sale":r.actual_sale or 0,"deposit":r.deposit or 0} for r in rows]

@router.get("/stats/vendor")
def vendor_stats(year: Optional[int]=None, month: Optional[int]=None, db: Session=Depends(get_db)):
    q = db.query(Reservation.vendor,func.count(Reservation.id).label("count"),func.sum(Reservation.headcount).label("headcount"),func.sum(Reservation.actual_sale).label("actual_sale"))
    if year: q = q.filter(extract("year",Reservation.event_date)==year)
    if month: q = q.filter(extract("month",Reservation.event_date)==month)
    rows = q.group_by(Reservation.vendor).order_by(func.sum(Reservation.actual_sale).desc()).all()
    return [{"vendor":r.vendor or "미지정","count":r.count,"headcount":r.headcount or 0,"actual_sale":r.actual_sale or 0} for r in rows]

@router.get("/vendors")
def get_vendors(db: Session=Depends(get_db)):
    rows = db.query(Reservation.vendor).distinct().all()
    return [r.vendor for r in rows if r.vendor]

@router.get("/years")
def get_years(db: Session=Depends(get_db)):
    rows = db.query(extract("year", Reservation.event_date).label("year")).distinct().order_by("year").all()
    return [int(r.year) for r in rows]

@router.get("/export/excel")
def export_excel(year: Optional[int]=None, month: Optional[int]=None, vendor: Optional[str]=None, db: Session=Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    q = db.query(Reservation)
    if year: q = q.filter(extract("year",Reservation.event_date)==year)
    if month: q = q.filter(extract("month",Reservation.event_date)==month)
    if vendor: q = q.filter(Reservation.vendor==vendor)
    rows = q.order_by(Reservation.event_date).all()
    wb = Workbook(); ws = wb.active; ws.title = "예약내역"
    hf = PatternFill("solid",start_color="2E4057"); hfont = Font(bold=True,color="FFFFFF",size=10)
    thin = Side(style="thin",color="CCCCCC"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
    headers = ["행사일","단체명","행사명","업체","판매가","DC후판매가","인원","DC금액","입금액","실판매가","전화번호","은행","계좌번호","예금주","담당자","메모"]
    widths = [12,16,16,10,12,12,8,12,12,12,14,10,16,10,10,20]
    for i,(h,w) in enumerate(zip(headers,widths),1):
        c=ws.cell(row=1,column=i,value=h); c.font=hfont; c.fill=hf; c.alignment=Alignment(horizontal="center"); c.border=border
        ws.column_dimensions[get_column_letter(i)].width=w
    for ri,r in enumerate(rows,2):
        vals=[r.event_date,r.group_name,r.event_name,r.vendor,r.sale_price,r.dc_sale_price,r.headcount,r.dc_amount,r.deposit,r.actual_sale,r.phone,r.bank,r.account_number,r.account_holder,r.manager,r.memo]
        for ci,v in enumerate(vals,1):
            c=ws.cell(row=ri,column=ci,value=v); c.border=border
            if ci in [5,6,8,9,10]: c.number_format="#,##0"
    tr=len(rows)+2; ws.cell(row=tr,column=1,value="합 계").font=Font(bold=True)
    for col in [7,8,9,10]:
        c = ws.cell(row=tr,column=col,value=f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{tr-1})")
        c.font = Font(bold=True)
        c.number_format = "#,##0"
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    from urllib.parse import quote
    fn=f"예약내역_{year or 'ALL'}{'_'+str(month)+'월' if month else ''}.xlsx"
    encoded=quote(fn)
    return StreamingResponse(buf,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f"attachment; filename*=UTF-8''{encoded}"})
