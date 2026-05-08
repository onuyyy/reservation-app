from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from urllib.parse import quote
import io
from database import get_db, CalendarEvent

router = APIRouter(prefix="/api/event-sheet", tags=["event-sheet"])

@router.get("/export/excel")
def export_excel(date: str, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from datetime import datetime

    rows = db.query(CalendarEvent).filter(
        CalendarEvent.event_date == date,
        CalendarEvent.is_confirmed == True
    ).order_by(CalendarEvent.id.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "행사표"

    dt = datetime.strptime(date, "%Y-%m-%d")
    title_text = f"{dt.month}월{dt.day}일"

    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = title_text
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    hfill = PatternFill("solid", start_color="2E4057")
    hfont = Font(bold=True, color="FFFFFF", size=10)
    halign = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["NO", "행사장", "단체명", "예상인원", "전화", "담당자", "식사", "차량", "비고"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=ci, value=h)
        c.font = hfont; c.fill = hfill; c.alignment = halign; c.border = border
    ws.row_dimensions[2].height = 22

    for ri, ev in enumerate(rows, 1):
        row_idx = ri + 2
        vals = [ri, ev.venue, ev.client, ev.headcount, ev.phone, ev.manager, ev.meal, ev.vehicle, ev.note]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=row_idx, column=ci, value=v)
            c.border = border
            c.alignment = Alignment(vertical="center")
        ws.row_dimensions[row_idx].height = 18

    last_data_row = len(rows) + 2
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="합계").font = Font(bold=True)
    sum_cell = ws.cell(row=total_row, column=4,
                       value=f"=SUM(D3:D{last_data_row})" if rows else 0)
    sum_cell.font = Font(bold=True)

    col_widths = [6, 14, 18, 8, 16, 10, 10, 10, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fn = f"행사표_{date}.xlsx"
    encoded = quote(fn)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"}
    )
